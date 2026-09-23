from docx import Document
from mcm_solarcheck.reporting.docx import render_docx
from mcm_solarcheck.reporting.model import InspectionReportModel,ReportEvidence

def test_docx_renderer_keeps_missing_temperature_explicit(tmp_path):
    model=InspectionReportModel('P','PV Thermal Inspection Report','Plant','review_complete','Celsius evidence is incomplete or unvalidated; no temperature claim may be inferred from raw values.',(ReportEvidence('F-1','thermal_anomaly_candidate',None,'T-1',None,20000,500.0,None,51.0,6.5,'Inspector','Checked'),),('Radiometric Celsius conversion is not validated.',))
    path=render_docx(model,tmp_path/'report.docx');assert path.exists()
    doc=Document(path);text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert 'F-1' in text;assert '20000' in text;assert 'Temperature [°C]' in text;assert 'Not available' in text
    assert 'Radiometric Celsius conversion is not validated.' in text


def test_docx_renderer_names_module_and_keeps_gps_optional(tmp_path):
    evidence=ReportEvidence('F-2','thermal_anomaly_candidate','M-0042','T-2',None,21000,600.0,None,None,None,'Inspector','Replace module')
    model=InspectionReportModel('P','PV Thermal Inspection Report','Plant','review_complete','No calibrated temperature claim.',(evidence,),())
    path=render_docx(model,tmp_path/'module-report.docx')
    doc=Document(path);text='\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert 'M-0042' in text
    assert 'Latitude' in text and 'Longitude' in text
    assert text.count('Not available')>=3
