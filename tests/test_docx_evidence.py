from pathlib import Path
from PIL import Image
from docx import Document
from mcm_solarcheck.reporting.docx import render_docx
from mcm_solarcheck.reporting.evidence_render import RenderedEvidence
from mcm_solarcheck.reporting.model import InspectionReportModel,ReportEvidence


def model():
    evidence=ReportEvidence('F-1','thermal_anomaly','M-1','T-1','V-1',20000,500,None,None,None,'expert','confirmed')
    return InspectionReportModel('P','PV Thermal Inspection Report','Project','review_complete','Celsius evidence unavailable.',(evidence,),())


def image(path):Image.new('RGB',(200,200),(100,120,140)).save(path)


def document_text(path):
    doc=Document(path);parts=[p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:parts.extend(p.text for p in cell.paragraphs)
    return '\n'.join(parts)


def test_docx_embeds_thermal_and_validated_rgb_evidence(tmp_path):
    thermal=tmp_path/'thermal.jpg';rgb=tmp_path/'rgb.jpg';image(thermal);image(rgb)
    out=render_docx(model(),tmp_path/'report.docx',rendered_evidence=(RenderedEvidence('F-1',thermal,rgb),))
    doc=Document(out);assert len(doc.inline_shapes)==2
    text=document_text(out);assert 'Thermal evidence' in text;assert 'RGB evidence — validated mapping' in text


def test_docx_never_claims_rgb_marker_when_rgb_evidence_missing(tmp_path):
    thermal=tmp_path/'thermal.jpg';image(thermal)
    out=render_docx(model(),tmp_path/'report.docx',rendered_evidence=(RenderedEvidence('F-1',thermal,None),))
    doc=Document(out);assert len(doc.inline_shapes)==1
    assert 'RGB evidence with validated marker unavailable.' in document_text(out)


def test_docx_remains_compatible_without_rendered_evidence(tmp_path):
    out=render_docx(model(),tmp_path/'report.docx')
    assert out.exists();assert 'Visual evidence: not rendered or unavailable.' in document_text(out)
