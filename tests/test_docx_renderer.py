from datetime import datetime, timezone
from docx import Document
from mcm_solarcheck.reporting.docx_renderer import REPORT_STANDARD_WORDING, render_docx
from mcm_solarcheck.reporting.report_model import InspectionReport, IrradianceSummary, OperatorSnapshot, ModuleReportDetail, ReportImage


def test_docx_renderer_keeps_customer_report_semantics(tmp_path):
    report=InspectionReport("REP-9","P1","Customer GmbH","Solarpark",datetime(2026,9,24,12,tzinfo=timezone.utc),"Inspector",850,12,5,site_address="Solarweg 1, 12345 Ort",irradiance=IrradianceSummary(750,"on-site sensor",700,810),release_status="reviewed")
    target=render_docx(report,tmp_path/"report.docx")
    assert target.is_file()
    doc=Document(target)
    text="\n".join(p.text for p in doc.paragraphs)+"\n"+"\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert REPORT_STANDARD_WORDING in text
    assert "Customer GmbH" in text and "Solarpark" in text
    assert "850" in text and "12" in text and "5" in text
    assert "Quelle: on-site sensor" in text
    assert "vollständig normkonform" not in text.lower()


def test_docx_renderer_does_not_invent_irradiance(tmp_path):
    report=InspectionReport("R","P","Customer","Site",datetime(2026,9,24,12,tzinfo=timezone.utc),"Inspector",10,0,0)
    target=render_docx(report,tmp_path/"report.docx")
    text="\n".join(p.text for p in Document(target).paragraphs)
    assert "Einstrahlung: nicht dokumentiert." in text


def test_docx_cover_uses_operator_snapshot(tmp_path):
    op=OperatorSnapshot("Operator GmbH","Werkstr. 1, 12345 Ort","office@example.invalid",phone="+49 123",website="https://example.invalid")
    report=InspectionReport("R","P","Customer","Site",datetime(2026,9,24,12,tzinfo=timezone.utc),"Inspector",10,0,0,operator=op)
    text="\n".join(p.text for p in Document(render_docx(report,tmp_path/"operator.docx")).paragraphs)
    assert "Operator GmbH" in text and "Werkstr. 1, 12345 Ort" in text
    assert "office@example.invalid" in text and "+49 123" in text


def test_docx_detail_marks_missing_images_explicitly(tmp_path):
    detail=ModuleReportDetail("M1","hotspot","confirmed",thermal_image=ReportImage("T1",str(tmp_path/"missing.png"),"thermal"))
    report=InspectionReport("R","P","Customer","Site",datetime(2026,9,24,12,tzinfo=timezone.utc),"Inspector",10,1,0,details=(detail,))
    doc=Document(render_docx(report,tmp_path/"missing-image.docx"))
    text="\n".join(p.text for p in doc.paragraphs)+"\n"+"\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert text.count("Bild nicht verfügbar")>=2
