"""DOCX renderer for the neutral inspection report model."""
from __future__ import annotations
from pathlib import Path
from docx import Document
from .model import InspectionReportModel

def render_docx(model:InspectionReportModel,output_path:str|Path)->Path:
    """Render a conservative report: missing evidence stays explicitly missing."""
    output=Path(output_path);output.parent.mkdir(parents=True,exist_ok=True)
    doc=Document();doc.add_heading(model.title,0);doc.add_paragraph(model.project_name)
    doc.add_heading('Inspection status',level=1);doc.add_paragraph(model.inspection_status)
    doc.add_heading('Thermal measurement status',level=1);doc.add_paragraph(model.temperature_statement)
    if model.warnings:
        doc.add_heading('Warnings',level=1)
        for warning in model.warnings:doc.add_paragraph(warning,style='List Bullet')
    doc.add_heading('Confirmed findings',level=1)
    if not model.evidence:doc.add_paragraph('No confirmed findings are available for this report.')
    for index,item in enumerate(model.evidence,1):
        doc.add_heading(f'Finding {index}: {item.finding_id}',level=2)
        table=doc.add_table(rows=0,cols=2)
        fields=(('Classification',item.classification),('Module',item.module_id),('Thermal frame',item.thermal_frame_id),('RGB frame',item.rgb_frame_id),('Raw sensor value',item.raw_value),('Raw delta from median',item.raw_delta_from_median),('Temperature [°C]',item.temperature_c),('Latitude',item.latitude),('Longitude',item.longitude),('Reviewer',item.reviewer),('Review note',item.review_note))
        for label,value in fields:
            cells=table.add_row().cells;cells[0].text=label;cells[1].text='Not available' if value is None else str(value)
    doc.save(output);return output
