"""Editable DOCX renderer for the neutral Phase 9 inspection report contract."""
from __future__ import annotations
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Mm, Pt
from .report_model import InspectionReport
from .presentation import detail_presentation, irradiance_text


REPORT_STANDARD_WORDING="Prüfbericht – Aufbau unter Berücksichtigung der DIN IEC/TS 62446-3 (VDE V 0126-23-3):2018-04"


def _value(value: object | None) -> str:
    return "—" if value is None or value=="" else str(value)


def render_docx(report: InspectionReport, destination: str | Path, *, banner_path: str | Path | None=None) -> Path:
    """Render an editable customer report without changing evidence semantics."""
    destination=Path(destination)
    document=Document()
    section=document.sections[0]
    section.page_width=Mm(210); section.page_height=Mm(297)
    section.top_margin=Mm(16); section.bottom_margin=Mm(16)
    section.left_margin=Mm(18); section.right_margin=Mm(18)
    normal=document.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(9)

    if banner_path is None and report.operator and report.operator.logo_path: banner_path=report.operator.logo_path
    if banner_path is not None:
        banner=Path(banner_path)
        if not banner.is_file(): raise FileNotFoundError(banner)
        document.add_picture(str(banner),width=Mm(174))
    document.add_heading(report.operator.company_name if report.operator else "MCM-SolarCheck",0)
    if report.operator:
        document.add_paragraph(report.operator.address)
        document.add_paragraph(" | ".join(v for v in (report.operator.email,report.operator.phone,report.operator.website) if v))
    document.add_paragraph(REPORT_STANDARD_WORDING)
    document.add_paragraph(f"Bericht: {report.report_id}")
    document.add_paragraph(f"Kunde: {report.customer_name}")
    document.add_paragraph(f"Anlage: {report.site_name}")
    document.add_paragraph(f"Standort: {_value(report.site_address)}")
    document.add_page_break()

    document.add_heading("Übersicht",level=1)
    table=document.add_table(rows=0,cols=2)
    for label,value in (
        ("Kunde",report.customer_name),("Kundenkontakt",report.customer_contact),("Kundenanschrift",report.customer_address),
        ("Kunden-E-Mail",report.customer_email),("Kundentelefon",report.customer_phone),("Kundenreferenz",report.customer_reference),("Auftragsreferenz",report.order_reference),("Anlage",report.site_name),
        ("Prüfbeginn",report.inspection_started_at.isoformat()),("Prüfer",report.inspector),
        ("PV-Module geprüft",report.total_modules),("Module mit dokumentiertem Befund",report.conspicuous_modules),
        ("Manuelle Prüfung erforderlich",report.manual_review_modules),
        ("Ohne dokumentierten Befund",report.modules_without_documented_finding),
    ):
        cells=table.add_row().cells; cells[0].text=label; cells[1].text=_value(value)
    document.add_paragraph(irradiance_text(report))
    for image,label in ((report.overview_rgb,"RGB-Übersicht"),(report.overview_thermal,"Thermal-Übersicht")):
        if image and Path(image.path).is_file():
            document.add_paragraph(label)
            document.add_picture(image.path,width=Mm(82))

    document.add_heading("Detailbefunde",level=1)
    if not report.details: document.add_paragraph("Keine freigegebenen Detailbefunde.")
    for detail in report.details:
        view=detail_presentation(detail)
        document.add_heading(view.heading,level=2)
        document.add_paragraph(f"Befund: {view.finding} | Review: {view.review}")
        if view.temperature: document.add_paragraph(view.temperature)
        if view.manual_inspection: document.add_paragraph(view.manual_inspection)
        images=document.add_table(rows=1,cols=2).cells
        for cell,img,label in ((images[0],detail.rgb_image,"RGB"),(images[1],detail.thermal_image,"Thermal")):
            context=view.rgb_context if label=="RGB" else view.thermal_context
            cell.text=label+(f"\n{context}" if context else "")
            if img and Path(img.path).is_file():
                cell.paragraphs[0].add_run().add_picture(img.path,width=Mm(76))
            else:
                cell.add_paragraph("Bild nicht verfügbar")

    document.add_heading("Zusammenfassung und Freigabe",level=1)
    document.add_paragraph(f"Prüfer: {report.inspector}")
    document.add_paragraph(f"Freigabestatus: {report.release_status}")
    for item in report.equipment:
        document.add_paragraph(f"Prüfmittel: {item.name}; ID: {_value(item.identifier)}; Kalibrierreferenz: {_value(item.calibration_reference)}")
    if report.provenance:
        document.add_paragraph(f"Software: {report.provenance.software_version}")
        document.add_paragraph(f"Datenprovenienz: {report.provenance.evidence_statement}")
    destination.parent.mkdir(parents=True,exist_ok=True)
    document.save(destination)
    return destination
